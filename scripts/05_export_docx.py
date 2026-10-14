"""
Step 5: Word(.docx) Export
- Step 1~4 의 산출물(metadata/summary/번역/이미지/캡션)을 기반으로
  글과 이미지가 자연스럽게 어우러진 Word 보고서를 생성한다.
- pandoc 같은 외부 바이너리 없이 순수 Python(python-docx)으로 동작.
- 결과: /final_report/battery_paper_analysis.docx

Usage:
    python scripts/05_export_docx.py
"""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
TEXT_DIR = ROOT / "output" / "text"
TRANSLATION_DIR = TEXT_DIR / "translations"
REPORT_DIR = ROOT / "final_report"
DOCX_PATH = REPORT_DIR / "battery_paper_analysis.docx"

MARKER_PATTERN = re.compile(r"\[\s*Figure\s+(\d+)\s+위치\s*\]", re.IGNORECASE)

SECTION_ORDER = [
    ("Abstract",     "1. Abstract (요약)"),
    ("Introduction", "2. Introduction (서론)"),
    ("Experimental", "3. Experimental (실험 방법)"),
    ("Results",      "4. Results and Discussion (결과 및 고찰)"),
    ("Conclusion",   "5. Conclusion (결론)"),
]

KOREAN_FONT = "Malgun Gothic"
IMAGE_WIDTH_IN = 5.5  # A4/Letter 기준 가독성이 좋은 폭


# ---------------------------------------------------------------------------
# Utilities
# ---------------------------------------------------------------------------

def _load_json(path: Path, default):
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def _load_caption(fig_num: int) -> str:
    p = TEXT_DIR / f"Figure_{fig_num}_caption.txt"
    if not p.exists():
        return ""
    return p.read_text(encoding="utf-8").strip()


def _apply_korean_font(run, *, size: int | None = None, bold: bool | None = None,
                       italic: bool | None = None, color: RGBColor | None = None) -> None:
    """한글/영문 모두 동일 폰트가 적용되도록 eastAsia 속성을 직접 설정."""
    run.font.name = KOREAN_FONT
    rpr = run._element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rpr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), KOREAN_FONT)
    rFonts.set(qn("w:ascii"), KOREAN_FONT)
    rFonts.set(qn("w:hAnsi"), KOREAN_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _configure_default_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = KOREAN_FONT
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rpr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), KOREAN_FONT)
    rFonts.set(qn("w:ascii"), KOREAN_FONT)
    rFonts.set(qn("w:hAnsi"), KOREAN_FONT)


# ---------------------------------------------------------------------------
# Block builders
# ---------------------------------------------------------------------------

def _add_rich_paragraph(doc: Document, text: str, *, size: int = 11,
                        indent: Inches | None = None, alignment=None):
    """**bold** 만 최소한으로 해석하는 단락 추가기."""
    p = doc.add_paragraph()
    if alignment is not None:
        p.alignment = alignment
    if indent is not None:
        p.paragraph_format.left_indent = indent
    parts = re.split(r"\*\*(.+?)\*\*", text)
    for i, part in enumerate(parts):
        if not part:
            continue
        run = p.add_run(part)
        _apply_korean_font(run, size=size, bold=(i % 2 == 1))
    return p


def _add_heading(doc: Document, text: str, level: int):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        _apply_korean_font(run, bold=True)
    return h


def _add_title_block(doc: Document, metadata: dict) -> None:
    translated = metadata.get("translated_title") or "[논문 제목 번역본]"
    original = metadata.get("original_title") or "[Original Title]"
    authors = metadata.get("authors") or "[Authors]"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(translated)
    _apply_korean_font(run, size=20, bold=True)

    for label, value in (("원제", original), ("저자", authors)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"{label}: ")
        _apply_korean_font(run, size=11, bold=True)
        run = p.add_run(value)
        _apply_korean_font(run, size=11, italic=True)

    doc.add_paragraph()  # spacer


def _add_summary_block(doc: Document, summary: dict) -> None:
    _add_heading(doc, "💡 3줄 핵심 요약", level=1)
    lines = summary.get("summary_3lines") or []
    while len(lines) < 3:
        lines.append("(요약을 생성하지 못했습니다.)")
    for i, item in enumerate(lines[:3], start=1):
        _add_rich_paragraph(doc, f"{i}. {item}")

    _add_heading(doc, "🎯 연구의 핵심 기여점", level=1)
    contrib = summary.get("contribution") or ["(핵심 기여점을 생성하지 못했습니다.)"]
    for item in contrib:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        _apply_korean_font(run, size=11)


def _add_image_block(doc: Document, image_path: Path, fig_num: int, caption: str) -> None:
    # 이미지 (페이지 중앙)
    if image_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        try:
            run.add_picture(str(image_path), width=Inches(IMAGE_WIDTH_IN))
        except Exception as exc:
            fallback = p.add_run(f"[Figure {fig_num} 이미지를 불러올 수 없습니다: {exc}]")
            _apply_korean_font(fallback, italic=True, color=RGBColor(0xC0, 0x39, 0x2B))
    else:
        _add_missing_warning(doc, fig_num)
        return

    # Figure 번호 라벨
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Figure {fig_num}")
    _apply_korean_font(run, size=10, bold=True)

    # AI 심층 해설 블록 (들여쓰기 + 이탤릭)
    if caption:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.right_indent = Inches(0.3)
        run = p.add_run("📊 AI Figure 심층 분석: ")
        _apply_korean_font(run, bold=True, size=11)
        run = p.add_run(caption)
        _apply_korean_font(run, italic=True, size=11)
    doc.add_paragraph()  # 블록 뒤 여백


def _add_missing_warning(doc: Document, fig_num: int) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    run = p.add_run(
        f"⚠️ [System Warning] Figure_{fig_num} 이미지를 원본 PDF에서 자동으로 "
        f"추출하지 못했습니다. 원본 논문을 참조해 주세요."
    )
    _apply_korean_font(run, bold=True, color=RGBColor(0xC0, 0x39, 0x2B))


def _emit_text_block(doc: Document, text: str) -> None:
    if not text or not text.strip():
        return
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        # 내부 마크다운 헤더(##)는 섹션 분할이 이미 끝났으므로 평문화
        para = re.sub(r"^#{1,6}\s*", "", para, flags=re.MULTILINE)
        _add_rich_paragraph(doc, para)


def _add_section_body(doc: Document, body: str, figure_mapping: dict) -> None:
    pos = 0
    for m in MARKER_PATTERN.finditer(body):
        _emit_text_block(doc, body[pos:m.start()])
        fig_num = int(m.group(1))
        meta = figure_mapping.get(f"Figure_{fig_num}")
        if meta and meta.get("status") == "extracted":
            image_path = ROOT / meta["file"]
            _add_image_block(doc, image_path, fig_num, _load_caption(fig_num))
        else:
            _add_missing_warning(doc, fig_num)
        pos = m.end()
    _emit_text_block(doc, body[pos:])


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def run() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)

    metadata = _load_json(TEXT_DIR / "metadata.json", {})
    summary = _load_json(TEXT_DIR / "summary.json", {"summary_3lines": [], "contribution": []})
    figure_mapping = _load_json(TEXT_DIR / "figure_mapping.json", {})

    doc = Document()
    _configure_default_style(doc)

    _add_title_block(doc, metadata)
    _add_summary_block(doc, summary)

    used_nums: set[int] = set()
    for canonical, header_text in SECTION_ORDER:
        body_path = TRANSLATION_DIR / f"{canonical}.md"
        if not body_path.exists():
            continue
        body = body_path.read_text(encoding="utf-8").strip()
        if not body:
            continue
        _add_heading(doc, header_text, level=1)
        _add_section_body(doc, body, figure_mapping)
        used_nums.update(int(m.group(1)) for m in MARKER_PATTERN.finditer(body))

    # 섹션에 포함되지 못한 추출 이미지는 Appendix 에서 노출
    leftovers = [
        meta for key, meta in figure_mapping.items()
        if meta.get("status") == "extracted" and meta.get("figure_number") not in used_nums
    ]
    if leftovers:
        _add_heading(doc, "📎 Appendix — 추가 Figure", level=1)
        for meta in sorted(leftovers, key=lambda m: m["figure_number"]):
            num = meta["figure_number"]
            image_path = ROOT / meta["file"]
            _add_image_block(doc, image_path, num, _load_caption(num))

    doc.save(str(DOCX_PATH))
    print("=" * 60)
    print(f"[DONE] Word 보고서: {DOCX_PATH}")
    print(f"[INFO] 이미지는 페이지 폭 {IMAGE_WIDTH_IN} in 로 맞춰 본문 사이에 삽입되었습니다.")


if __name__ == "__main__":
    run()
